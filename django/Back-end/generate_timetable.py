import os
import django
import random
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'Back-end.settings')
django.setup()

from api.models import Teachers, Classes

def _header_style(cell):
    cell.font = Font(bold=True, color="FFFFFFFF")
    cell.fill = PatternFill("solid", fgColor="FF1F2937")
    cell.alignment = Alignment(horizontal="center", vertical="center")

teachers = list(Teachers.objects.select_related('user').all())
classes = list(Classes.objects.all())

if not teachers or not classes:
    print("Not enough data in DB to generate timetable.")
    exit(1)

wb = Workbook()
ws = wb.active
ws.title = "Emploi"

headers = ["jour", "heure-debut", "heure-fin", "matiere", "professeur", "classe", "salle"]
ws.append(headers)
for cell in ws[1]:
    _header_style(cell)

jours = ["Lundi", "Mardi", "Mercredi", "jeudi", "Vendredi"]
salles = ["Amphi A", "B101", "B102", "Lab1", "Lab2", "Lab3", "Lab4", "B201", "B202"]
matieres = ["Algorithmique", "Bases de données", "Génie Logiciel", "Réseaux", "Mathématiques", "Anglais", "Systèmes d'exploitation"]

slots = [
    ("08:30", "10:00"),
    ("10:15", "11:45"),
    ("14:00", "15:30"),
    ("15:45", "17:15")
]

# Trackers to avoid overlaps
used_teachers = set()  # (jour, time_start, teacher_name)
used_rooms = set()     # (jour, time_start, room_name)

for c in classes:
    c_label = f"{c.niveau}-{c.classe_section}-{c.classe_num}"
    
    # 3 random days
    chosen_days = random.sample(jours, 3)
    
    for jour in chosen_days:
        # Assign 1 to 2 random slots on this day
        chosen_slots = random.sample(slots, random.randint(1, 2))
        
        for time_start, time_end in chosen_slots:
            # Find an available teacher
            available_teachers = [t for t in teachers if (jour, time_start, t.user.username) not in used_teachers]
            # Find an available room
            available_rooms = [r for r in salles if (jour, time_start, r) not in used_rooms]
            
            if available_teachers and available_rooms:
                teacher = random.choice(available_teachers)
                t_name = teacher.user.username
                salle = random.choice(available_rooms)
                matiere = random.choice(matieres)
                
                # Mark as used
                used_teachers.add((jour, time_start, t_name))
                used_rooms.add((jour, time_start, salle))
                
                ws.append([jour, time_start, time_end, matiere, t_name, c_label, salle])

filepath = "/home/nouri/mini-projet-issat-pfe-app/timetable_generated.xlsx"
wb.save(filepath)
print(f"Successfully generated timetable with {ws.max_row - 1} slots without any collisions at {filepath}")
