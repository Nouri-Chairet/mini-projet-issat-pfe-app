"""
Generators for the ready-to-use Excel templates that admins download
before filling in real data.

All templates are produced in-memory with openpyxl and returned as
:class:`django.http.HttpResponse` instances with the correct
Content-Disposition header.
"""
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _header_style(cell):
    cell.font = Font(bold=True, color="FFFFFFFF")
    cell.fill = PatternFill("solid", fgColor="FF1F2937")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 48)


def _xlsx_response(workbook: Workbook, filename: str) -> HttpResponse:
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type=XLSX_MIME)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def build_timetable_template() -> HttpResponse:
    """
    Build a clean timetable Excel template matching the columns the
    backend parser expects (`_parse_schedule_dataframe` in
    `admin_panel/views.py`).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Emploi"

    headers = [
        "jour",
        "heure-debut",
        "heure-fin",
        "matiere",
        "professeur",
        "classe",
        "salle",
    ]
    ws.append(headers)
    for cell in ws[1]:
        _header_style(cell)

    sample_rows = [
        ["Lundi", "08:00", "10:00", "Algorithmique", "Ahmed Trabelsi", "1-Prépa-1", "B102"],
        ["Lundi", "10:15", "12:15", "Analyse", "Samia Ben Salah", "1-Prépa-1", "B102"],
        ["Mardi", "08:00", "10:00", "Programmation", "Mehdi Karoui", "2-L-LSI-3", "Lab1"],
        ["Mercredi", "13:00", "15:00", "Bases de données", "Hela Mahjoub", "2-L-LSI-3", "Lab2"],
        ["jeudi", "08:00", "10:00", "Réseaux", "Karim Jendoubi", "2-L-LSI-3", "B205"],
        ["Vendredi", "08:00", "10:00", "Anglais", "Sonia Dhaouadi", "1-Prépa-1", "B101"],
    ]
    for row in sample_rows:
        ws.append(row)
    _autosize(ws)

    help_ws = wb.create_sheet("Help")
    help_ws.append(["Column", "Description", "Example / Allowed values"])
    for cell in help_ws[1]:
        _header_style(cell)
    help_rows = [
        ["jour", "Day of the week (French)", "Lundi, Mardi, Mercredi, jeudi, Vendredi, Samedi"],
        ["heure-debut", "Start time, 24h format", "08:00"],
        ["heure-fin", "End time, 24h format (must be after heure-debut)", "10:00"],
        ["matiere", "Subject name", "Algorithmique"],
        ["professeur", "Teacher username (must already exist in the system)", "Ahmed Trabelsi"],
        ["classe", "Class label '<niveau>-<section>-<num>'", "1-Prépa-1, 2-L-LSI-3, 3-L-Mécanique-1, 1-Cycle Ingénieur-2"],
        ["salle", "Room label", "B102, Lab1, Amphi A"],
    ]
    for row in help_rows:
        help_ws.append(row)
    _autosize(help_ws)

    return _xlsx_response(wb, "timetable_template.xlsx")


def build_pfe_bulk_template() -> HttpResponse:
    """
    Build the bulk PFE assignment Excel template.

    Each row maps one student to one supervisor with the chosen PFE
    title. Students that already have a PFE are rejected by the
    importer's dry-run, so the template only documents the columns.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "PFE"

    headers = [
        "title",
        "student_email",
        "student_name",
        "supervisor_email",
        "supervisor_name",
        "description",
    ]
    ws.append(headers)
    for cell in ws[1]:
        _header_style(cell)

    sample_rows = [
        [
            "Agentic AI for student support",
            "student1@issat.example",
            "Yassine Khaldi",
            "teacher1@issat.example",
            "Ahmed Trabelsi",
            "Build an AI agent that answers student admin questions.",
        ],
        [
            "Smart timetable optimizer",
            "student2@issat.example",
            "Imen Bouazizi",
            "teacher2@issat.example",
            "Samia Ben Salah",
            "Optimization-based weekly schedule planner.",
        ],
        [
            "Predictive maintenance dashboard",
            "student3@issat.example",
            "Mohamed Ali Saidi",
            "teacher3@issat.example",
            "Mehdi Karoui",
            "",
        ],
    ]
    for row in sample_rows:
        ws.append(row)
    _autosize(ws)

    help_ws = wb.create_sheet("Help")
    help_ws.append(["Column", "Required", "Notes"])
    for cell in help_ws[1]:
        _header_style(cell)
    help_rows = [
        ["title", "yes", "PFE subject title."],
        ["student_email", "yes (or student_name)", "Email of an existing student account. Used as primary identifier."],
        ["student_name", "fallback", "Used to look up the student if email is empty."],
        ["supervisor_email", "yes (or supervisor_name)", "Email of an existing teacher account."],
        ["supervisor_name", "fallback", "Used to look up the teacher if email is empty."],
        ["description", "no", "Optional description of the PFE subject."],
        ["—", "—", "Each student can be assigned only ONE PFE subject. Duplicates are rejected."],
    ]
    for row in help_rows:
        help_ws.append(row)
    _autosize(help_ws)

    return _xlsx_response(wb, "pfe_bulk_template.xlsx")
